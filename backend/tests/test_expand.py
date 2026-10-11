"""Run with: python -m unittest discover -s backend/tests (python-pptx required)."""

import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from unittest.mock import patch
import shutil

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.errors import error_code
from app.expand import A, P, Target, UnsupportedAnimationError, _align_table_highlights, _highlight_band, _inherit_placeholder_autofit, _mask_formula_fallbacks, _normalize_symbol_fonts, _paragraph_bands, _prefer_formula_fallbacks, _slide_order, _state, _xml, expand_pptx


def make_fixture(path: Path, unsupported=False, connector=False):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(.3), Inches(5), Inches(.5))
    title.text = "Baseline"
    ids = []
    for number in range(1, 4):
        if connector and number == 1:
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(1), Inches(4), Inches(1))
        else:
            shape = slide.shapes.add_textbox(Inches(1), Inches(number), Inches(5), Inches(.5))
            shape.text = f"步骤 {number}"
        ids.append(shape.shape_id)
    timing = etree.SubElement(slide._element, f"{{{P}}}timing")
    root_list = etree.SubElement(timing, f"{{{P}}}tnLst")
    root_par = etree.SubElement(root_list, f"{{{P}}}par")
    root_time = etree.SubElement(root_par, f"{{{P}}}cTn", id="1", dur="indefinite", nodeType="tmRoot")
    children = etree.SubElement(root_time, f"{{{P}}}childTnLst")
    for offset, sid in enumerate(ids, 2):
        par = etree.SubElement(children, f"{{{P}}}par")
        effect = etree.SubElement(par, f"{{{P}}}cTn", id=str(offset), nodeType="clickEffect",
                                  presetClass="exit" if unsupported else "entr")
        effect_children = etree.SubElement(effect, f"{{{P}}}childTnLst")
        setting = etree.SubElement(effect_children, f"{{{P}}}set")
        behavior = etree.SubElement(setting, f"{{{P}}}cBhvr")
        etree.SubElement(behavior, f"{{{P}}}cTn", id=str(offset + 10), dur="1")
        target = etree.SubElement(behavior, f"{{{P}}}tgtEl")
        etree.SubElement(target, f"{{{P}}}spTgt", spid=str(sid))
        names = etree.SubElement(behavior, f"{{{P}}}attrNameLst")
        etree.SubElement(names, f"{{{P}}}attrName").text = "style.visibility"
        etree.SubElement(setting, f"{{{P}}}to").append(etree.Element(f"{{{P}}}strVal", val="visible"))
    deck.save(path)


class ExpansionTest(unittest.TestCase):
    def test_upload_limit_is_20_mib(self):
        from app import main, storage
        with tempfile.TemporaryDirectory() as directory, patch.object(storage, "DATA_DIR", Path(directory)):
            with TestClient(main.app) as client:
                self.assertEqual(client.get("/api/health").json()["max_file_size"], 20 * 1024 * 1024)
                accepted = client.post("/api/uploads", json={"filename": "limit.pptx", "size": storage.MAX_FILE})
                self.assertEqual(accepted.status_code, 200, accepted.text)
                rejected = client.post("/api/uploads", json={"filename": "too-large.pptx", "size": storage.MAX_FILE + 1})
                self.assertEqual(rejected.status_code, 400, rejected.text)
                self.assertEqual(rejected.headers["x-error-code"], "invalidUpload")

    def test_click_entrances_create_each_state(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.pptx", Path(directory) / "expanded.pptx"
            make_fixture(source)
            report = expand_pptx(source, output)
            self.assertEqual((report.original_slides, report.output_pages, report.expanded_slides), (1, 4, 1))
            with ZipFile(output) as archive:
                paths = _slide_order(archive)
                for page, expected in enumerate((0, 1, 2, 3), 1):
                    slide = etree.fromstring(archive.read(paths[page - 1]))
                    texts = slide.findall(f".//{{{A}}}t")
                    self.assertEqual([text.text for text in texts], ["Baseline"] + [f"步骤 {i}" for i in range(1, expected + 1)])

    def test_connector_entrance_keeps_cumulative_states(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.pptx", Path(directory) / "expanded.pptx"
            make_fixture(source, connector=True)
            report = expand_pptx(source, output)
            self.assertEqual(report.output_pages, 4)
            self.assertEqual(report.fallback_slides, [])
            with ZipFile(output) as archive:
                connectors = [len(etree.fromstring(archive.read(part)).findall(f".//{{{P}}}cxnSp"))
                              for part in _slide_order(archive)]
            self.assertEqual(connectors, [0, 1, 1, 1])

    def test_unsupported_effect_identifies_slide_instead_of_exporting_wrong_order(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.pptx", Path(directory) / "expanded.pptx"
            make_fixture(source, unsupported=True)
            with self.assertRaises(UnsupportedAnimationError) as failure:
                expand_pptx(source, output)
            self.assertEqual(failure.exception.slide_number, 1)
            self.assertEqual(error_code(str(failure.exception)), "unsupportedAnimation")

    def test_embedded_formula_artwork_is_masked_by_paragraph_without_changing_visible_pixels(self):
        image = Image.new("RGBA", (40, 52), (10, 20, 30, 0))
        for top, bottom in [(2, 5), (16, 19), (24, 27), (38, 41), (46, 49)]:
            for y in range(top, bottom + 1):
                for x in range(4, 36):
                    image.putpixel((x, y), (80 + y, 12, 140, 255))
        stream = BytesIO()
        image.save(stream, format="PNG")
        _, groups = _paragraph_bands(stream.getvalue(), 3)
        self.assertEqual(groups, [(2, 5), (16, 27), (38, 49)])
        state = _xml(b'''<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
            xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main"
            xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
            <mc:AlternateContent><mc:Choice><p:sp><p:nvSpPr><p:cNvPr id="4"/></p:nvSpPr>
            <p:txBody><a:p/><a:p><a14:m/></a:p><a:p/></p:txBody></p:sp></mc:Choice>
            <mc:Fallback><p:pic><p:blipFill><a:blip r:embed="rId4"/></p:blipFill></p:pic></mc:Fallback>
            </mc:AlternateContent></p:sld>''')
        rels = _xml(b'''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
            <Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
            Target="../media/image.png"/></Relationships>''')
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.zip"
            with ZipFile(path, "w") as archive:
                archive.writestr("ppt/media/image.png", stream.getvalue())
            with ZipFile(path) as archive:
                media = _mask_formula_fallbacks(state, rels, archive, "ppt/slides/slide1.xml",
                                                [Target("4", (0, 1))], 1)
        self.assertEqual(len(media), 1)
        masked = Image.open(BytesIO(media[0][1])).convert("RGBA")
        self.assertEqual(masked.getpixel((10, 3))[3], 0)
        self.assertEqual(masked.getpixel((10, 25))[3], 0)
        self.assertEqual(masked.getpixel((10, 40)), image.getpixel((10, 40)))
        _prefer_formula_fallbacks(state, state)
        self.assertFalse(state.xpath('.//*[local-name()="AlternateContent"]'))
        self.assertEqual(state.xpath('.//*[local-name()="blip"]/@r:embed',
                                     namespaces={"r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}),
                         ["rIdAnimImage1_1"])

    def test_formula_overlay_waits_for_animated_colored_background(self):
        source = _xml(b'''<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
            xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"
            xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
            <p:cSld><p:spTree>
              <p:sp><p:nvSpPr><p:cNvPr id="7"/></p:nvSpPr><p:spPr>
                <a:xfrm><a:off x="0" y="0"/><a:ext cx="1000" cy="500"/></a:xfrm>
                <a:solidFill><a:srgbClr val="8C1511"/></a:solidFill>
              </p:spPr></p:sp>
              <mc:AlternateContent><mc:Choice><p:sp>
                <p:nvSpPr><p:cNvPr id="8"/></p:nvSpPr><p:spPr>
                  <a:xfrm><a:off x="100" y="100"/><a:ext cx="800" cy="300"/></a:xfrm>
                </p:spPr><p:txBody><a:p>
                  <a:r><a:rPr><a:solidFill><a:schemeClr val="bg1"/></a:solidFill></a:rPr><a:t>Formula: </a:t></a:r>
                  <m:oMath><m:r><a:rPr><a:solidFill><a:schemeClr val="bg1"/></a:solidFill></a:rPr><m:t>x</m:t></m:r></m:oMath>
                </a:p></p:txBody>
              </p:sp></mc:Choice><mc:Fallback><p:pic><p:blipFill>
                <a:blip r:embed="rId1"/>
              </p:blipFill></p:pic></mc:Fallback></mc:AlternateContent>
            </p:spTree></p:cSld>
        </p:sld>''')
        before = _xml(_state(source, [[Target("7")]], 0))
        after = _xml(_state(source, [[Target("7")]], 1))
        self.assertEqual(before.xpath('.//p:cNvPr/@id', namespaces={"p": P}), [])
        self.assertEqual(after.xpath('.//p:cNvPr/@id', namespaces={"p": P}), ["7", "8"])
        source.find(f"{{{P}}}cSld/{{{P}}}spTree").remove(source.xpath('.//p:sp[p:nvSpPr/p:cNvPr[@id="7"]]', namespaces={"p": P})[0])
        without_backing = _xml(_state(source, [], 0))
        self.assertEqual(without_backing.xpath('.//p:cNvPr/@id', namespaces={"p": P}), [])

    def test_wingdings_arrow_uses_a_portable_symbol_font(self):
        slide = _xml(b'''<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <a:r><a:rPr><a:sym typeface="Wingdings" charset="2"/></a:rPr><a:t>\xef\x83\xa0</a:t></a:r>
            <a:r><a:rPr><a:latin typeface="Calibri"/></a:rPr><a:t>Other</a:t></a:r>
        </p:sld>''')
        _normalize_symbol_fonts(slide)
        self.assertEqual(slide.xpath('.//a:r[1]/a:t/text()', namespaces={"a": A}), ["\u2794"])
        self.assertEqual(slide.xpath('.//a:r[1]/a:rPr/a:latin/@typeface', namespaces={"a": A}),
                         ["Noto Sans Symbols2"])
        self.assertEqual(slide.xpath('.//a:r[2]/a:t/text()', namespaces={"a": A}), ["Other"])

    def test_placeholder_inherits_layout_autofit_without_overriding_explicit_choice(self):
        layout = _xml(b'''<p:sldLayout xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree>
            <p:sp><p:nvSpPr><p:nvPr><p:ph type="title" idx="0"/></p:nvPr></p:nvSpPr>
              <p:txBody><a:bodyPr><a:normAutofit/></a:bodyPr></p:txBody></p:sp>
            </p:spTree></p:cSld></p:sldLayout>''')
        slide = _xml(b'''<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree>
            <p:sp><p:nvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>
              <p:txBody><a:bodyPr/><a:p><a:r><a:t>Long title</a:t></a:r></a:p></p:txBody></p:sp>
            </p:spTree></p:cSld></p:sld>''')
        _inherit_placeholder_autofit(slide, [layout])
        self.assertEqual(len(slide.xpath('.//p:sp/p:txBody/a:bodyPr/a:normAutofit',
                                         namespaces={"p": P, "a": A})), 1)
        body = slide.xpath('.//p:sp/p:txBody/a:bodyPr', namespaces={"p": P, "a": A})[0]
        body.clear()
        etree.SubElement(body, f"{{{A}}}noAutofit")
        _inherit_placeholder_autofit(slide, [layout])
        self.assertEqual([etree.QName(item).localname for item in body], ["noAutofit"])

    def test_full_table_formula_snapshot_is_used_once(self):
        slide = _xml(b'''<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
            xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"
            xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
            <p:cSld><p:spTree><mc:AlternateContent><mc:Choice><p:graphicFrame>
              <p:nvGraphicFramePr><p:cNvPr id="5" name="Table"/></p:nvGraphicFramePr>
              <p:xfrm><a:off x="10" y="20"/><a:ext cx="100" cy="60"/></p:xfrm>
              <a:graphic><a:graphicData><a:tbl><a:tr><a:tc><m:oMath/></a:tc></a:tr></a:tbl></a:graphicData></a:graphic>
            </p:graphicFrame></mc:Choice><mc:Fallback><p:graphicFrame><a:graphic><a:graphicData>
              <a:tbl><a:tr><a:tc><a:blip r:embed="rId3"/></a:tc>
              <a:tc><a:blip r:embed="rId3"/></a:tc></a:tr></a:tbl>
            </a:graphicData></a:graphic></p:graphicFrame></mc:Fallback></mc:AlternateContent>
            </p:spTree></p:cSld></p:sld>''')
        snapshots = _prefer_formula_fallbacks(slide, slide)
        self.assertEqual(snapshots, [("5", "rId3")])
        self.assertEqual(len(slide.xpath('.//p:pic', namespaces={"p": P})), 1)
        self.assertEqual(slide.xpath('.//p:pic/p:spPr/a:xfrm/a:ext/@cy',
                                     namespaces={"p": P, "a": A}), ["60"])
        image = Image.new("RGB", (100, 100), "white")
        for y in range(75, 95):
            for x in range(100):
                image.putpixel((x, y), (243, 224, 168))
        self.assertEqual(_highlight_band(image), (75, 95))

    def test_table_highlight_outline_hugs_the_yellow_row(self):
        image = Image.new("RGB", (100, 100), "white")
        for y in range(75, 95):
            for x in range(100):
                image.putpixel((x, y), (243, 224, 168))
        stream = BytesIO()
        image.save(stream, format="PNG")
        slide = _xml(b'''<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
            <p:cSld><p:spTree>
              <p:pic><p:nvPicPr><p:cNvPr id="5"/></p:nvPicPr>
                <p:blipFill><a:blip r:embed="rId3"/></p:blipFill>
                <p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1000" cy="1000"/></a:xfrm></p:spPr>
              </p:pic>
              <p:sp><p:spPr><a:xfrm><a:off x="0" y="700"/><a:ext cx="1000" cy="300"/></a:xfrm>
                <a:noFill/><a:ln/></p:spPr></p:sp>
            </p:spTree></p:cSld></p:sld>''')
        rels = _xml(b'''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
            <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
            Target="../media/table.png"/></Relationships>''')
        with tempfile.TemporaryDirectory() as directory:
            archive_path = Path(directory) / "source.pptx"
            with ZipFile(archive_path, "w") as archive:
                archive.writestr("ppt/media/table.png", stream.getvalue())
            with ZipFile(archive_path) as archive:
                _align_table_highlights(slide, rels, archive, "ppt/slides/slide1.xml", [("5", "rId3")])
        outline = slide.xpath('.//p:sp/p:spPr/a:xfrm', namespaces={"p": P, "a": A})[0]
        self.assertEqual(outline.find(f"{{{A}}}off").get("y"), "740")
        self.assertEqual(outline.find(f"{{{A}}}ext").get("cy"), "220")

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice is not installed")
    def test_upload_convert_and_download(self):
        from app import main, storage, worker
        with tempfile.TemporaryDirectory() as directory:
            with patch.object(storage, "DATA_DIR", Path(directory)), \
                 patch("app.main.redis.Redis.from_url") as connection, \
                 patch("app.main.set_status"), \
                 patch.object(main.convert, "apply_async"), \
                 patch("app.worker.set_status") as worker_status:
                connection.return_value.set.return_value = True
                source = Path(directory) / "fixture.pptx"
                make_fixture(source)
                data = source.read_bytes()
                with TestClient(main.app) as client:
                    created = client.post("/api/uploads", json={"filename": "fixture.pptx", "size": len(data)})
                    self.assertEqual(created.status_code, 200, created.text)
                    upload_id = created.json()["upload_id"]
                    chunk = client.put(f"/api/uploads/{upload_id}/chunks/0", content=data)
                    self.assertEqual(chunk.status_code, 200, chunk.text)
                    self.assertEqual(client.get(f"/api/uploads/{upload_id}").json()["uploaded_chunks"], [0])
                    completed = client.post(f"/api/uploads/{upload_id}/complete")
                    self.assertEqual(completed.status_code, 200, completed.text)
                    worker.convert(upload_id)
                    self.assertEqual(worker_status.call_args.kwargs["message_key"], "complete")
                    self.assertEqual(worker_status.call_args.kwargs["fallback_slides"], [])
                    pdf = client.get(f"/api/jobs/{upload_id}/download")
                    self.assertEqual(pdf.status_code, 200, pdf.text[:200])
                    self.assertIn("animation-steps.pdf", pdf.headers["content-disposition"])
                    preview = client.get(f"/api/jobs/{upload_id}/preview")
                    self.assertEqual(preview.status_code, 200)
                    self.assertEqual(preview.headers["content-type"], "application/pdf")
                    self.assertTrue(preview.headers["content-disposition"].startswith("inline;"))
                    self.assertEqual(preview.headers["cache-control"], "private, no-store")
                    self.assertEqual(preview.content, pdf.content)
                    first_bytes = client.get(f"/api/jobs/{upload_id}/preview", headers={"Range": "bytes=0-15"})
                    self.assertEqual(first_bytes.status_code, 206)
                    self.assertEqual(first_bytes.content, pdf.content[:16])
                    image_preview = client.get(f"/api/jobs/{upload_id}/preview/pages/2")
                    self.assertEqual(image_preview.status_code, 200)
                    self.assertEqual(image_preview.headers["content-type"], "image/png")
                    self.assertEqual(image_preview.headers["cache-control"], "private, no-store")
                    with Image.open(BytesIO(image_preview.content)) as image:
                        self.assertGreater(image.width, 1000)
                    invalid_page = client.get(f"/api/jobs/{upload_id}/preview/pages/5")
                    self.assertEqual(invalid_page.status_code, 404)
                    self.assertEqual(invalid_page.headers["x-error-code"], "invalidPreviewPage")
                    traditional = client.get(f"/api/jobs/{upload_id}/download?lang=zh-TW")
                    self.assertIn("%E5%8B%95%E7%95%AB", traditional.headers["content-disposition"])
                    result = Path(directory) / "result.pdf"
                    result.write_bytes(pdf.content)
                    pages = PdfReader(str(result)).pages
                    self.assertEqual(len(pages), 4)
                    for index, page in enumerate(pages):
                        content = page.extract_text()
                        self.assertIn("Baseline", content)
                        for step in range(1, 4):
                            self.assertEqual(f" {step}" in content, step <= index)


if __name__ == "__main__":
    unittest.main()
